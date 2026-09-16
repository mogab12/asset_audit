"""Leitura e gravação de planilhas no formato exportado pelo Effort.

Observação: o Effort exporta arquivos com extensão .xls que na verdade são
.xlsx. Por isso o formato é detectado pelo conteúdo, não pela extensão.
"""
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from app.utils import sem_acentos, texto

ASSINATURA_XLSX = b"PK"
ASSINATURA_XLS = b"\xD0\xCF\x11\xE0"


class ErroPlanilha(Exception):
    pass


def _normalizar_titulo(titulo):
    return " ".join(sem_acentos(titulo).lower().split())


_TITULO_PARA_CAMPO = {
    _normalizar_titulo(titulo): campo for campo, titulo in config.COLUNAS.items()
}


def _linhas_xlsx(conteudo):
    livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    aba = livro.active
    for linha in aba.iter_rows(values_only=True):
        yield list(linha)
    livro.close()


def _linhas_xls(conteudo):
    try:
        import xlrd
    except ImportError as erro:
        raise ErroPlanilha("Para ler .xls antigos instale o pacote xlrd.") from erro
    livro = xlrd.open_workbook(file_contents=conteudo)
    aba = livro.sheet_by_index(0)
    for i in range(aba.nrows):
        yield aba.row_values(i)


def ler(conteudo):
    """Recebe os bytes do arquivo e devolve uma lista de dicionários."""
    if conteudo.startswith(ASSINATURA_XLSX):
        linhas = _linhas_xlsx(conteudo)
    elif conteudo.startswith(ASSINATURA_XLS):
        linhas = _linhas_xls(conteudo)
    else:
        raise ErroPlanilha("O arquivo não parece ser uma planilha do Excel.")

    try:
        cabecalho = next(linhas)
    except StopIteration:
        raise ErroPlanilha("A planilha está vazia.")

    mapa = {}  # índice da coluna -> campo interno
    for indice, titulo in enumerate(cabecalho):
        campo = _TITULO_PARA_CAMPO.get(_normalizar_titulo(titulo))
        if campo:
            mapa[indice] = campo

    campos_encontrados = set(mapa.values())
    if "equipamento" not in campos_encontrados or not (
            {"setor_origem", "setor_atual"} & campos_encontrados):
        raise ErroPlanilha(
            "Colunas obrigatórias não encontradas. A planilha precisa ter "
            "pelo menos 'Equipamento' e 'Setor Origem' ou 'Setor Atual'.")

    itens = []
    for linha in linhas:
        item = {campo: texto(linha[i]) if i < len(linha) else ""
                for i, campo in mapa.items()}
        if any(item.values()):
            itens.append(item)
    return itens


def gerar(itens):
    """Gera um .xlsx com as mesmas colunas da importação. Devolve BytesIO."""
    livro = Workbook()
    aba = livro.active
    aba.title = config.NOME_ABA_EXPORTACAO
    campos = list(config.COLUNAS.keys())

    fonte = Font(name="Arial", size=10)
    fonte_titulo = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fundo_titulo = PatternFill("solid", fgColor="1F4E5A")

    aba.append([config.COLUNAS[c] for c in campos])
    for celula in aba[1]:
        celula.font = fonte_titulo
        celula.fill = fundo_titulo
        celula.alignment = Alignment(vertical="center")

    for item in itens:
        # Tudo como texto para não perder zeros à esquerda (ex.: 062.013158)
        aba.append([texto(item.get(c)) or None for c in campos])

    larguras = {"equipamento": 60, "observacao": 35, "setor_origem": 28,
                "setor_atual": 28, "executante": 30, "motivo": 26}
    for indice, campo in enumerate(campos, start=1):
        letra = get_column_letter(indice)
        aba.column_dimensions[letra].width = larguras.get(campo, 18)
        for celula in aba[letra][1:]:
            celula.font = fonte
            celula.number_format = "@"

    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions

    saida = io.BytesIO()
    livro.save(saida)
    saida.seek(0)
    return saida
